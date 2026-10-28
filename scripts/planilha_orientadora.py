"""Planilha para a orientadora preencher: pontos da Reitoria e população/área.

Uso (a partir da raiz do projeto, com o venv do backend ativo e o Postgres local no ar):

    python scripts/planilha_orientadora.py gerar     # cria Docs/EcoCampus_dados_para_preencher.xlsx
    python scripts/planilha_orientadora.py previa    # lê a planilha preenchida e mostra o que mudaria

`previa` NÃO grava nada: só compara a planilha com o banco e imprime as diferenças.
Os dois comandos recusam bancos que não sejam locais.
"""
from __future__ import annotations

import argparse
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from dotenv import load_dotenv  # noqa: E402

load_dotenv(ROOT / "backend" / ".env")

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
from sqlalchemy import func, select  # noqa: E402

from app.db.models import Campus, Lancamento, PontoMedicao  # noqa: E402
from app.db.session import SessionLocal, engine  # noqa: E402

ARQUIVO_PADRAO = ROOT / "Docs" / "EcoCampus_dados_para_preencher.xlsx"

ABA_PONTOS = "Pontos da Reitoria"
ABA_CAMPI = "População e área"
REITORIA_ID = 6

OPCAO_REITORIA = "Reitoria (está certo)"
OPCAO_OUTRA = "Outra (escrever em Observação)"

# A linha 1 é a instrução, a 2 é o cabeçalho; os dados começam na 3.
LINHA_CABECALHO = 2
PRIMEIRA_LINHA = 3

COLUNAS_PONTOS = [
    "Tipo",
    "Endereço da fatura",
    "Identificador",
    "Local/descrição",
    "Unidade correta",
    "Observação",
    "Código interno (não alterar)",
]
COLUNAS_CAMPI = [
    "Campus",
    "População total (alunos + docentes + técnicos administrativos)",
    "Área construída (m²)",
    "Data de referência",
    "Fonte/setor que forneceu",
]

FONTE = "Arial"
AMARELO = PatternFill("solid", start_color="FFF2CC")
CINZA = PatternFill("solid", start_color="E7E6E6")
BORDA = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def _exige_banco_local() -> None:
    host = engine.url.host or ""
    if host not in {"localhost", "127.0.0.1"}:
        raise SystemExit(f"Este script só roda contra o banco local (host atual: {host!r}).")


def _identificador(ponto: PontoMedicao) -> str:
    if ponto.tipo.value == "agua":
        return f"RGI/PDE: {ponto.identificador}" if ponto.identificador else ""
    partes = []
    if ponto.identificador:
        partes.append(f"Nº da UC: {ponto.identificador}")
    if ponto.codigo_instalacao:
        partes.append(f"Instalação: {ponto.codigo_instalacao}")
    return " | ".join(partes)


def _sem_acento(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")


def _ajusta_larguras(ws, maximo: int = 60) -> None:
    for indice, coluna in enumerate(ws.iter_cols(min_row=LINHA_CABECALHO), start=1):
        maior = max((len(str(c.value)) for c in coluna if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(indice)].width = min(max(maior + 3, 12), maximo)


def _monta_folha(ws, instrucao: str, colunas: list[str], editaveis: set[int]) -> None:
    ultima = get_column_letter(len(colunas))
    ws.merge_cells(f"A1:{ultima}1")
    ws["A1"] = instrucao
    ws["A1"].font = Font(name=FONTE, size=11, bold=True)
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 62

    for indice, titulo in enumerate(colunas, start=1):
        celula = ws.cell(row=LINHA_CABECALHO, column=indice, value=titulo)
        celula.font = Font(name=FONTE, bold=True)
        celula.fill = AMARELO if indice in editaveis else CINZA
        celula.alignment = Alignment(wrap_text=True, vertical="center")
        celula.border = BORDA
    ws.row_dimensions[LINHA_CABECALHO].height = 32
    # Instrução e cabeçalho ficam visíveis ao rolar.
    # Referência em texto: `ws.cell(...)` criaria a célula e empurraria os dados uma linha.
    ws.freeze_panes = f"A{PRIMEIRA_LINHA}"


def _formata_dados(ws, n_colunas: int, editaveis: set[int]) -> None:
    for linha in ws.iter_rows(min_row=PRIMEIRA_LINHA, max_col=n_colunas):
        for celula in linha:
            celula.font = Font(name=FONTE)
            celula.border = BORDA
            celula.alignment = Alignment(vertical="center", wrap_text=celula.column in editaveis)
            if celula.column in editaveis:
                celula.fill = AMARELO


def gerar(destino: Path) -> None:
    _exige_banco_local()
    with SessionLocal() as db:
        campi = db.scalars(select(Campus).order_by(Campus.id)).all()
        pontos = db.scalars(select(PontoMedicao).where(PontoMedicao.campus_id == REITORIA_ID)).all()

    pontos.sort(key=lambda p: (p.tipo.value, _sem_acento(p.endereco or "").upper(), p.id))
    opcoes = [c.nome for c in campi if c.id != REITORIA_ID] + [OPCAO_REITORIA, OPCAO_OUTRA]
    lista = ",".join(opcoes)
    # Lista embutida na validação: o Excel aceita até 255 caracteres.
    assert len(lista) <= 255 and not any("," in o for o in opcoes)

    wb = Workbook()

    # --- Aba 1 -----------------------------------------------------------
    ws = wb.active
    ws.title = ABA_PONTOS
    _monta_folha(
        ws,
        "Indique a qual unidade cada conta pertence de fato.\n"
        "Preencha apenas as colunas em amarelo: em \"Unidade correta\", clique na célula e escolha uma opção da lista. "
        "Se a conta é mesmo da Reitoria, escolha \"Reitoria (está certo)\". "
        "Se a unidade não está na lista, escolha \"Outra\" e escreva qual é na coluna Observação.",
        COLUNAS_PONTOS,
        editaveis={5, 6},
    )
    for ponto in pontos:
        ws.append(
            [
                "Água" if ponto.tipo.value == "agua" else "Energia",
                ponto.endereco or "",
                _identificador(ponto),
                ponto.nome,
                None,
                None,
                ponto.id,
            ]
        )
    _formata_dados(ws, len(COLUNAS_PONTOS), editaveis={5, 6})
    ultima = PRIMEIRA_LINHA + len(pontos) - 1

    validacao = DataValidation(type="list", formula1=f'"{lista}"', allow_blank=True, showDropDown=False)
    validacao.errorTitle = "Opção inválida"
    validacao.error = "Escolha uma das opções da lista. Se a unidade não estiver na lista, escolha \"Outra\" e explique na coluna Observação."
    validacao.promptTitle = "Unidade correta"
    validacao.prompt = "Clique na seta e escolha a unidade a que esta conta pertence."
    validacao.showErrorMessage = True
    validacao.showInputMessage = True
    ws.add_data_validation(validacao)
    validacao.add(f"E{PRIMEIRA_LINHA}:E{ultima}")

    _ajusta_larguras(ws)
    ws.column_dimensions["E"].width = 34
    ws.column_dimensions["F"].width = 45
    ws.column_dimensions["G"].hidden = True

    # --- Aba 2 -----------------------------------------------------------
    ws2 = wb.create_sheet(ABA_CAMPI)
    _monta_folha(
        ws2,
        "Informe a população e a área construída de cada campus. Preencha apenas as colunas em amarelo.\n"
        "População e área: só números, sem ponto de milhar. Data de referência: mês e ano a que os números se referem.\n"
        "Exemplo de linha preenchida:  2150  |  12400  |  03/2026  |  Pró-Reitoria de Administração",
        COLUNAS_CAMPI,
        editaveis={2, 3, 4, 5},
    )
    for campus in campi:
        ws2.append([campus.nome, None, None, None, None])
    _formata_dados(ws2, len(COLUNAS_CAMPI), editaveis={2, 3, 4, 5})
    ultima2 = PRIMEIRA_LINHA + len(campi) - 1

    inteiro = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    inteiro.errorTitle = "Valor inválido"
    inteiro.error = "Digite só o número de pessoas, sem ponto nem texto. Exemplo: 2150"
    inteiro.showErrorMessage = True
    decimal = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    decimal.errorTitle = "Valor inválido"
    decimal.error = "Digite só o número de metros quadrados, sem ponto de milhar nem texto. Exemplo: 12400"
    decimal.showErrorMessage = True
    ws2.add_data_validation(inteiro)
    ws2.add_data_validation(decimal)
    inteiro.add(f"B{PRIMEIRA_LINHA}:B{ultima2}")
    decimal.add(f"C{PRIMEIRA_LINHA}:C{ultima2}")
    for linha in range(PRIMEIRA_LINHA, ultima2 + 1):
        ws2.cell(row=linha, column=2).number_format = "#,##0"
        ws2.cell(row=linha, column=3).number_format = "#,##0.##"
        ws2.cell(row=linha, column=4).number_format = "@"  # texto: "03/2026" não vira data

    _ajusta_larguras(ws2)
    ws2.column_dimensions["B"].width = 36
    ws2.column_dimensions["E"].width = 40

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    print(f"Planilha gerada: {destino}")
    print(f"  {ABA_PONTOS}: {len(pontos)} pontos | {ABA_CAMPI}: {len(campi)} campi")
    print(f"  Opções da lista: {opcoes}")


# --- Prévia ---------------------------------------------------------------


def _texto(valor) -> str:
    return str(valor).strip() if valor is not None else ""


def _numero(valor) -> tuple[float | None, str | None]:
    """Converte o que foi digitado em número. Retorna (valor, erro)."""
    if valor is None or _texto(valor) == "":
        return None, None
    if isinstance(valor, (int, float)):
        return float(valor), None
    bruto = _texto(valor).replace("m²", "").replace("m2", "").strip()
    # Planilha em português: ponto é milhar, vírgula é decimal.
    limpo = bruto.replace(".", "").replace(" ", "").replace(",", ".")
    try:
        return float(limpo), None
    except ValueError:
        return None, f"não é número: {valor!r}"


def previa(arquivo: Path) -> None:
    _exige_banco_local()
    if not arquivo.exists():
        raise SystemExit(f"Arquivo não encontrado: {arquivo}")
    wb = load_workbook(arquivo, data_only=True)
    for aba in (ABA_PONTOS, ABA_CAMPI):
        if aba not in wb.sheetnames:
            raise SystemExit(f"A planilha não tem a aba {aba!r}.")

    with SessionLocal() as db:
        campi = db.scalars(select(Campus).order_by(Campus.id)).all()
        pontos = {p.id: p for p in db.scalars(select(PontoMedicao)).all()}
        lancamentos_por_ponto = dict(
            db.execute(
                select(Lancamento.ponto_medicao_id, func.count())
                .where(Lancamento.ponto_medicao_id.isnot(None))
                .group_by(Lancamento.ponto_medicao_id)
            ).all()
        )
    id_por_nome = {c.nome: c.id for c in campi}
    nome_por_id = {c.id: c.nome for c in campi}

    print(f"PRÉVIA — nada será gravado.  Arquivo: {arquivo}\n")

    # --- Aba 1 -----------------------------------------------------------
    print(f"== {ABA_PONTOS} ==")
    mover, manter, outras, em_branco, problemas = [], [], [], [], []
    vistos: set[int] = set()
    for linha in wb[ABA_PONTOS].iter_rows(min_row=PRIMEIRA_LINHA, max_col=len(COLUNAS_PONTOS), values_only=True):
        tipo, endereco, ident, local, unidade, obs, ponto_id = linha
        if not any(v is not None and _texto(v) for v in linha):
            continue
        rotulo = f"{_texto(tipo)} · {_texto(local)} · {_texto(endereco)}"
        try:
            ponto = pontos.get(int(ponto_id))
        except (TypeError, ValueError):
            ponto = None
        if ponto is None:
            problemas.append(f"{rotulo}: código interno ausente ou desconhecido ({ponto_id!r})")
            continue
        if ponto.id in vistos:
            problemas.append(f"{rotulo}: ponto {ponto.id} aparece mais de uma vez")
            continue
        vistos.add(ponto.id)
        # A linha pode ter sido trocada de lugar: confere se o identificador ainda bate.
        if _identificador(ponto) != _texto(ident):
            problemas.append(f"{rotulo}: o identificador não confere com o ponto {ponto.id} do banco")
            continue

        unidade, obs = _texto(unidade), _texto(obs)
        if not unidade:
            em_branco.append(rotulo)
        elif unidade == OPCAO_REITORIA:
            manter.append(rotulo)
        elif unidade == OPCAO_OUTRA:
            outras.append(f"{rotulo}  →  observação: {obs or '(vazia!)'}")
        elif unidade in id_por_nome:
            destino = id_por_nome[unidade]
            if destino == ponto.campus_id:
                manter.append(rotulo)
            else:
                n = lancamentos_por_ponto.get(ponto.id, 0)
                mover.append(
                    f"ponto {ponto.id:>3}  {rotulo}\n"
                    f"            {nome_por_id[ponto.campus_id]}  →  {unidade}   ({n} lançamento(s) acompanham)"
                    + (f"\n            obs.: {obs}" if obs else "")
                )
        else:
            problemas.append(f"{rotulo}: unidade fora da lista: {unidade!r}")

    faltando = [p for p in pontos.values() if p.campus_id == REITORIA_ID and p.id not in vistos]

    def _bloco(titulo: str, itens: list[str]) -> None:
        print(f"\n{titulo} ({len(itens)})")
        for item in itens:
            print(f"  - {item}")

    _bloco("Seriam REATRIBUÍDOS", mover)
    _bloco("Ficam na Reitoria", manter)
    _bloco("\"Outra\" — precisam de decisão manual", outras)
    _bloco("Ainda em branco", em_branco)
    _bloco("Pontos da Reitoria que sumiram da planilha", [f"ponto {p.id}: {p.nome} · {p.endereco}" for p in faltando])
    _bloco("PROBLEMAS (linhas ignoradas)", problemas)

    # --- Aba 2 -----------------------------------------------------------
    print(f"\n== {ABA_CAMPI} ==")
    alteracoes, sem_dado, erros = [], [], []
    for linha in wb[ABA_CAMPI].iter_rows(min_row=PRIMEIRA_LINHA, max_col=len(COLUNAS_CAMPI), values_only=True):
        nome, populacao, area, data_ref, fonte = linha
        nome = _texto(nome)
        if not nome:
            continue
        if nome not in id_por_nome:
            erros.append(f"campus desconhecido: {nome!r}")
            continue
        campus = next(c for c in campi if c.nome == nome)
        pop, erro_pop = _numero(populacao)
        ar, erro_area = _numero(area)
        for erro in (erro_pop, erro_area):
            if erro:
                erros.append(f"{nome}: {erro}")
        if pop is not None and pop != int(pop):
            erros.append(f"{nome}: população não é número inteiro ({pop})")
            pop = None
        if pop is None and ar is None:
            sem_dado.append(nome)
            continue
        if isinstance(data_ref, (date, datetime)):
            data_ref = data_ref.strftime("%m/%Y")
        partes = []
        for rotulo_campo, atual, novo in (
            ("população", campus.populacao, None if pop is None else int(pop)),
            # A coluna do banco é inteira: a área é arredondada.
            ("área (m²)", campus.area_construida_m2, None if ar is None else round(ar)),
        ):
            if novo is None:
                partes.append(f"{rotulo_campo} não informada, fica {atual}")
            elif novo == atual:
                partes.append(f"{rotulo_campo} {atual} (sem mudança)")
            else:
                partes.append(f"{rotulo_campo} {atual} → {novo}")
        alteracoes.append(
            f"{nome}: " + "; ".join(partes) + f"   [ref. {_texto(data_ref) or '?'}; fonte: {_texto(fonte) or '?'}]"
        )

    _bloco("Seriam ATUALIZADOS", alteracoes)
    _bloco("Sem dado na planilha", sem_dado)
    _bloco("PROBLEMAS", erros)
    print("\nNada foi gravado.")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("comando", choices=["gerar", "previa"])
    parser.add_argument("--arquivo", type=Path, default=ARQUIVO_PADRAO)
    args = parser.parse_args()
    if args.comando == "gerar":
        gerar(args.arquivo)
    else:
        previa(args.arquivo)


if __name__ == "__main__":
    main()
